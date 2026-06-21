"""批量导入供应商 — 解析 Excel 文件并入库。"""

import io
import uuid

import openpyxl

from app.core.logging import get_logger
from app.db.mongo import get_db
from app.db.postgres import get_cursor
from app.services.embedding import encode_single

logger = get_logger()

# Expected column headers (Chinese)
COLUMN_MAP = {
    "企业全称": "name",
    "公司名称": "name",
    "供应商名称": "name",
    "name": "name",
    "统一社会信用代码": "unified_code",
    "信用代码": "unified_code",
    "unified_code": "unified_code",
    "主营品类": "categories",
    "品类": "categories",
    "主营产品": "categories",
    "categories": "categories",
    "经营地域": "regions",
    "地域": "regions",
    "区域": "regions",
    "regions": "regions",
    "状态": "status",
    "status": "status",
}


def import_suppliers_from_excel(file_content: bytes, filename: str) -> dict:
    """解析 Excel 文件，批量导入供应商。

    预期表格格式：
    - 首行为表头，列名需包含在 COLUMN_MAP 中
    - 每行一个供应商
    - 品类/地域支持逗号分隔

    返回：{"imported": N, "skipped": N, "errors": [...]}
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
    db = get_db()

    imported = 0
    skipped = 0
    errors: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # Parse header
        headers = [str(h).strip() if h else "" for h in rows[0]]
        col_map: dict[str, int] = {}
        for idx, h in enumerate(headers):
            if h in COLUMN_MAP:
                col_map[COLUMN_MAP[h]] = idx

        if "name" not in col_map:
            errors.append(f"Sheet '{sheet_name}' 缺少企业全称列（支持：企业全称/公司名称/供应商名称/name）")
            continue

        # Parse data rows
        for row_idx, row in enumerate(rows[1:], 2):
            try:
                name = _cell(row, col_map.get("name"))
                if not name:
                    skipped += 1
                    continue

                sid = str(uuid.uuid4())
                categories = _cell_list(row, col_map.get("categories"))
                regions = _cell_list(row, col_map.get("regions"))
                unified_code = _cell(row, col_map.get("unified_code"))
                status = _cell(row, col_map.get("status")) or "prospective"

                # Check for duplicate by name
                existing = db["suppliers"].find_one({"name": name})
                if existing:
                    db["suppliers"].update_one(
                        {"name": name},
                        {"$set": {
                            "categories": categories,
                            "regions": regions,
                            "unified_code": unified_code,
                            "status": status,
                            "source": "excel_import",
                            "embedding_dirty": False,
                        }},
                    )
                    skipped += 1
                    continue

                # Insert into MongoDB
                db["suppliers"].insert_one({
                    "_id": sid,
                    "name": name,
                    "unified_code": unified_code or None,
                    "categories": categories,
                    "regions": regions,
                    "status": status,
                    "source": "excel_import",
                    "embedding_dirty": False,
                })

                # Build PG vector
                content_parts = [name] + categories + regions
                embedding = encode_single(" ".join(content_parts))
                vec_str = "[" + ",".join(str(v) for v in embedding) + "]"
                with get_cursor() as (conn, cur):
                    cur.execute(
                        """INSERT INTO supplier_profiles (id, supplier_name, content, embedding, metadata)
                           VALUES (%s, %s, %s, %s::vector, %s)""",
                        (sid, name, " ".join(content_parts), vec_str, "{}"),
                    )

                imported += 1

            except Exception as e:
                errors.append(f"Sheet '{sheet_name}' 第{row_idx}行: {str(e)}")
                logger.error("supplier_import_error", sheet=sheet_name, row=row_idx, error=str(e))

    logger.info("supplier_import_done", filename=filename, imported=imported, skipped=skipped)
    return {"imported": imported, "skipped": skipped, "errors": errors}


def _cell(row, idx: int | None) -> str:
    if idx is None or idx >= len(row):
        return ""
    val = row[idx]
    return str(val).strip() if val is not None else ""


def _cell_list(row, idx: int | None) -> list[str]:
    raw = _cell(row, idx)
    if not raw:
        return []
    return [s.strip() for s in raw.replace("，", ",").split(",") if s.strip()]
