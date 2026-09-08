"""Import internal supplier-material relationships for sourcing recall."""

from __future__ import annotations

import hashlib
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
from pymongo import ReplaceOne

from app.core.logging import get_logger
from app.db.mongo import get_db

logger = get_logger(__name__)

SOURCE_PROVIDER = "internal_supplier_material_list"
RELATION_COLLECTION = "internal_supplier_material_relations"
BATCH_COLLECTION = "internal_supplier_material_import_batches"
EXPECTED_HEADERS = ("供应商代码", "供应商名称", "物料号", "物料名称", "基地")


def _text(value: object) -> str | None:
    value = str(value).strip() if value is not None else ""
    return value or None


def _relation_id(supplier_code: str, material_number: str, base_code: str) -> str:
    digest = hashlib.sha256(f"{supplier_code}\n{material_number}\n{base_code}".encode()).hexdigest()
    return f"{SOURCE_PROVIDER}:{digest}"


def parse_history_row(row: tuple[object, ...]) -> dict[str, Any] | None:
    """Convert an internal supplier-material-base row while retaining ID strings."""
    values = [_text(value) for value in row[:5]]
    if len(values) != 5 or any(value is None for value in values):
        return None
    supplier_code, supplier_name, material_number, material_name, base_code = values
    return {
        "_id": _relation_id(supplier_code, material_number, base_code),
        "supplier_code": supplier_code,
        "supplier_name": supplier_name,
        "material_number": material_number,
        "material_name": material_name,
        "base_code": base_code,
        "source_provider": SOURCE_PROVIDER,
    }


def import_internal_supplier_material_file(path: str | Path) -> dict[str, Any]:
    """Idempotently import a validated supplier-material relationship workbook."""
    source_path = Path(path).expanduser().resolve()
    if not source_path.is_file():
        raise ValueError(f"历史零件供应商文件不存在: {source_path}")

    workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
    if header != EXPECTED_HEADERS:
        raise ValueError(f"历史零件供应商文件字段不匹配，期望：{'、'.join(EXPECTED_HEADERS)}")

    db = get_db()
    batch_id = str(uuid.uuid4())
    started_at = datetime.now(timezone.utc)
    db[BATCH_COLLECTION].insert_one({
        "_id": batch_id,
        "source_provider": SOURCE_PROVIDER,
        "source_file": source_path.name,
        "source_path": str(source_path),
        "status": "running",
        "started_at": started_at,
    })

    operations: list[ReplaceOne] = []
    raw_rows = 0
    skipped_rows = 0
    now = datetime.now(timezone.utc)
    try:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            raw_rows += 1
            relation = parse_history_row(row)
            if not relation:
                skipped_rows += 1
                continue
            relation.update({"latest_batch_id": batch_id, "updated_at": now})
            operations.append(ReplaceOne({"_id": relation["_id"]}, relation, upsert=True))

        for index in range(0, len(operations), 1000):
            db[RELATION_COLLECTION].bulk_write(operations[index:index + 1000], ordered=False)
        db[RELATION_COLLECTION].create_index(
            [("material_number", 1), ("supplier_code", 1), ("base_code", 1)], unique=True
        )
        db[RELATION_COLLECTION].create_index([("material_name", 1), ("supplier_code", 1)])

        result = {
            "batch_id": batch_id,
            "source_file": source_path.name,
            "raw_rows": raw_rows,
            "skipped_rows": skipped_rows,
            "relation_operations": len(operations),
            "status": "completed",
            "started_at": started_at,
            "completed_at": datetime.now(timezone.utc),
        }
        db[BATCH_COLLECTION].update_one({"_id": batch_id}, {"$set": result})
        logger.info("internal_supplier_material_import_completed", **result)
        return result
    except Exception as exc:
        db[BATCH_COLLECTION].update_one(
            {"_id": batch_id},
            {"$set": {"status": "failed", "failed_at": datetime.now(timezone.utc), "error": str(exc)}},
        )
        logger.error("internal_supplier_material_import_failed", batch_id=batch_id, error=str(exc))
        raise
