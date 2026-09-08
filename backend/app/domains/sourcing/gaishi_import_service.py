"""Import manually obtained Gasgoo supplier exports into the external candidate library."""

from __future__ import annotations

import hashlib
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from pymongo import ReplaceOne

from app.core.logging import get_logger
from app.db.mongo import get_db

logger = get_logger(__name__)

SOURCE_PROVIDER = "gasgoo_manual_export"
PROFILE_COLLECTION = "gaishi_supplier_profiles"
LINK_COLLECTION = "gaishi_supplier_category_links"
BATCH_COLLECTION = "gaishi_import_batches"

EXPECTED_HEADERS = (
    "公司名称", "标签1", "标签2", "标签3", "企业性质", "收录时间", "所在地",
    "注册资金", "成立时间", "主营产品", "配套客户",
)
CHIP_HEADERS = ("企业名称", "地区", "上市公司", "所属类型")


def normalize_company_name(value: object) -> str:
    """Return a stable matching key without changing the displayed company name."""
    return "".join(str(value or "").split()).casefold()


def _text(value: object) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None


def _profile_id(company_key: str) -> str:
    return f"{SOURCE_PROVIDER}:{company_key}"


def _link_id(category: str, company_key: str) -> str:
    digest = hashlib.sha256(f"{SOURCE_PROVIDER}\n{category}\n{company_key}".encode()).hexdigest()
    return f"{SOURCE_PROVIDER}:{digest}"


def parse_gaishi_row(category: str, source_file: str, row: tuple[object, ...]) -> tuple[dict[str, Any], dict[str, Any]] | None:
    """Map one standard Gasgoo row to a profile and its category evidence link."""
    company_name = _text(row[0] if row else None)
    company_key = normalize_company_name(company_name)
    if not company_name or not company_key:
        return None

    tags = [tag for tag in (_text(row[index]) for index in (1, 2, 3)) if tag]
    profile_id = _profile_id(company_key)
    profile = {
        "_id": profile_id,
        "company_name": company_name,
        "company_key": company_key,
        "tags": tags,
        "enterprise_nature": _text(row[4]),
        "included_age": _text(row[5]),
        "location": _text(row[6]),
        "registered_capital": _text(row[7]),
        "established_date": _text(row[8]),
        "main_products": _text(row[9]),
        "supporting_customers": _text(row[10]),
        "source_provider": SOURCE_PROVIDER,
    }
    link = {
        "_id": _link_id(category, company_key),
        "supplier_profile_id": profile_id,
        "company_name": company_name,
        "company_key": company_key,
        "category": category,
        "source_file": source_file,
        "source_provider": SOURCE_PROVIDER,
    }
    return profile, link


def parse_gaishi_chip_row(source_file: str, row: tuple[object, ...]) -> tuple[dict[str, Any], dict[str, Any]] | None:
    """Map the separately exported chip candidate workbook into the same library."""
    company_name = _text(row[0] if row else None)
    company_key = normalize_company_name(company_name)
    if not company_name or not company_key:
        return None

    profile_id = _profile_id(company_key)
    profile = {
        "_id": profile_id,
        "company_name": company_name,
        "company_key": company_key,
        "tags": [tag for tag in (_text(row[2]), _text(row[3])) if tag],
        "enterprise_nature": _text(row[3]),
        "included_age": None,
        "location": _text(row[1]),
        "registered_capital": None,
        "established_date": None,
        "main_products": None,
        "supporting_customers": None,
        "source_provider": SOURCE_PROVIDER,
    }
    link = {
        "_id": _link_id("芯片", company_key),
        "supplier_profile_id": profile_id,
        "company_name": company_name,
        "company_key": company_key,
        "category": "芯片",
        "source_file": source_file,
        "source_provider": SOURCE_PROVIDER,
    }
    return profile, link


def _batched(items: Iterable[Any], size: int = 1000) -> Iterable[list[Any]]:
    batch: list[Any] = []
    for item in items:
        batch.append(item)
        if len(batch) >= size:
            yield batch
            batch = []
    if batch:
        yield batch


def import_gaishi_directory(directory: str | Path) -> dict[str, Any]:
    """Idempotently import standard Gasgoo category workbooks from a directory.

    This importer intentionally writes to an external-candidate library.  It never
    creates records in ``suppliers`` because a Gasgoo listing is not proof of
    internal supplier qualification or legal-entity identity.
    """
    root = Path(directory).expanduser().resolve()
    if not root.is_dir():
        raise ValueError(f"盖世数据目录不存在: {root}")

    files = sorted(root.glob("*.xlsx"))
    if not files:
        raise ValueError(f"盖世数据目录中没有 Excel 文件: {root}")

    db = get_db()
    batch_id = str(uuid.uuid4())
    started_at = datetime.now(timezone.utc)
    db[BATCH_COLLECTION].insert_one({
        "_id": batch_id,
        "source_provider": SOURCE_PROVIDER,
        "source_directory": str(root),
        "status": "running",
        "started_at": started_at,
    })

    profile_ops: list[ReplaceOne] = []
    link_ops: list[ReplaceOne] = []
    files_imported = 0
    raw_rows = 0
    skipped_rows = 0
    invalid_files: list[dict[str, str]] = []
    now = datetime.now(timezone.utc)

    try:
        for path in files:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            header = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
            parser = None
            if header == EXPECTED_HEADERS:
                parser = lambda row: parse_gaishi_row(path.stem, path.name, row)
            elif header == CHIP_HEADERS and path.stem == "芯片":
                parser = lambda row: parse_gaishi_chip_row(path.name, row)
            else:
                invalid_files.append({"file": path.name, "reason": "字段结构不属于盖世企业候选表"})
                continue

            files_imported += 1
            for row in sheet.iter_rows(min_row=2, values_only=True):
                raw_rows += 1
                parsed = parser(row)
                if not parsed:
                    skipped_rows += 1
                    continue
                profile, link = parsed
                profile.update({"latest_batch_id": batch_id, "updated_at": now})
                link.update({"latest_batch_id": batch_id, "updated_at": now})
                profile_ops.append(ReplaceOne({"_id": profile["_id"]}, profile, upsert=True))
                link_ops.append(ReplaceOne({"_id": link["_id"]}, link, upsert=True))

        for operations in _batched(profile_ops):
            db[PROFILE_COLLECTION].bulk_write(operations, ordered=False)
        for operations in _batched(link_ops):
            db[LINK_COLLECTION].bulk_write(operations, ordered=False)

        db[PROFILE_COLLECTION].create_index("company_key", unique=True)
        db[LINK_COLLECTION].create_index([("category", 1), ("company_key", 1)], unique=True)
        db[LINK_COLLECTION].create_index([("category", 1), ("supplier_profile_id", 1)])

        completed_at = datetime.now(timezone.utc)
        result = {
            "batch_id": batch_id,
            "source_directory": str(root),
            "files_found": len(files),
            "files_imported": files_imported,
            "raw_rows": raw_rows,
            "skipped_rows": skipped_rows,
            "profile_operations": len(profile_ops),
            "category_link_operations": len(link_ops),
            "invalid_files": invalid_files,
            "status": "completed",
            "started_at": started_at,
            "completed_at": completed_at,
        }
        db[BATCH_COLLECTION].update_one({"_id": batch_id}, {"$set": result})
        logger.info("gaishi_import_completed", **{key: value for key, value in result.items() if key != "invalid_files"})
        return result
    except Exception as exc:
        db[BATCH_COLLECTION].update_one(
            {"_id": batch_id},
            {"$set": {"status": "failed", "failed_at": datetime.now(timezone.utc), "error": str(exc)}},
        )
        logger.error("gaishi_import_failed", batch_id=batch_id, error=str(exc))
        raise
