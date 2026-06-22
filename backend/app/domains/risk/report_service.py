from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from app.domains.risk.repo_company import get_risk_info, get_risk_indicators
from app.domains.risk.repo_financial import get_financial_metrics
from app.domains.risk.company_service import get_company_profile


def _has_val(v) -> bool:
    """Return True if value is not None and not 0.0 (sentinel for missing data)."""
    return v is not None and v != 0.0


def generate_excel(company_name: str) -> bytes:
    profile = get_company_profile(company_name)
    risk = get_risk_info(company_name)
    indicators = get_risk_indicators(company_name)
    fin = get_financial_metrics(company_name)

    wb = Workbook()
    header_font = Font(bold=True, size=12)
    header_fill = PatternFill("solid", fgColor="F0F0EB")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    def _style_header(ws, row, cols):
        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border

    # --- Sheet 1: 概览 ---
    ws1 = wb.active
    ws1.title = "风险概览"
    ws1.cell(1, 1, f"企业风险评估报告 — {company_name}").font = Font(bold=True, size=14)
    ws1.merge_cells("A1:D1")
    ws1.cell(2, 1, f"法人代表: {profile.legal_person}  |  注册资本: {profile.registered_capital}  |  成立: {profile.establish_time}")
    ws1.merge_cells("A2:D2")

    ws1.cell(4, 1, "指标"); ws1.cell(4, 2, "数值")
    _style_header(ws1, 4, 2)
    rows = [
        ("诉讼数量", risk.lawsuit_count),
        ("被执行记录", indicators["executed_count"]),
        ("失信记录", indicators["dishonesty_count"]),
        ("重大诉讼", "是" if indicators["major_lawsuit"] else "否"),
        ("经营异常", risk.abnormal_operation_count),
        ("行政处罚", risk.administrative_penalty_count),
        ("法人频繁变更", "是" if indicators["legal_person_change_frequent"] else "否"),
        ("对外担保", indicators["guarantee_count"]),
        ("股权质押", indicators["pledge_count"]),
        ("破产/清算", indicators["bankruptcy_count"]),
        ("环保处罚", indicators["env_penalty_count"]),
    ]
    for i, (k, v) in enumerate(rows, 5):
        ws1.cell(i, 1, k).border = thin_border
        ws1.cell(i, 2, v).border = thin_border

    ws1.column_dimensions["A"].width = 18
    ws1.column_dimensions["B"].width = 14

    # --- Sheet 2: 财务指标 ---
    ws2 = wb.create_sheet("财务指标")
    ws2.cell(1, 1, "指标"); ws2.cell(1, 2, "数值"); ws2.cell(1, 3, "说明")
    _style_header(ws2, 1, 3)
    if fin:
        metrics = [
            ("营收增长率", f"{fin.revenue_growth * 100:.1f}%", "近一年营收同比"),
            ("净利增长率", f"{fin.net_profit_growth * 100:.1f}%", "近一年净利同比"),
            ("资产负债率", f"{fin.debt_ratio * 100:.1f}%", "总负债/总资产"),
            ("每股现金流", f"¥{fin.cash_flow:.2f}", "经营活动现金流"),
            ("ROE", f"{fin.roe * 100:.1f}%" if fin.roe else "—", "净资产收益率"),
            ("净利率", f"{fin.net_profit_margin * 100:.1f}%" if fin.net_profit_margin else "—", "净利润/营收"),
            ("流动比率", f"{fin.current_ratio:.2f}" if fin.current_ratio else "—", "流动资产/流动负债"),
            ("速动比率", f"{fin.quick_ratio:.2f}" if fin.quick_ratio else "—", "速动资产/流动负债"),
            ("存货周转率", f"{fin.inventory_turnover:.1f}" if _has_val(fin.inventory_turnover) else "—", ""),
            ("应收款周转天数", f"{fin.ar_turnover_days:.0f}天" if _has_val(fin.ar_turnover_days) else "—", ""),
            ("扣非利润占比", f"{fin.recurring_profit_ratio * 100:.1f}%" if _has_val(fin.recurring_profit_ratio) else "—", ""),
            ("产权比率", f"{fin.equity_ratio:.2f}" if _has_val(fin.equity_ratio) else "—", ""),
            ("营收趋势", "下滑" if fin.revenue_trend and fin.revenue_trend < 0 else "稳定", "近3年"),
            ("负债趋势", "上升" if fin.debt_trend and fin.debt_trend > 0 else "稳定", "近3年"),
            ("净利趋势", "下滑" if fin.net_profit_trend and fin.net_profit_trend < 0 else "稳定", "近3年"),
        ]
    else:
        metrics = [("—", "无财报数据", "")]
    for i, (k, v, note) in enumerate(metrics, 2):
        ws2.cell(i, 1, k).border = thin_border
        ws2.cell(i, 2, v).border = thin_border
        ws2.cell(i, 3, note).border = thin_border

    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 20

    # --- Sheet 3: 风险明细 ---
    ws3 = wb.create_sheet("风险明细")
    ws3.cell(1, 1, "类别"); ws3.cell(1, 2, "指标"); ws3.cell(1, 3, "数值")
    _style_header(ws3, 1, 3)
    row = 2
    for cat, label, val in [
        ("司法", "诉讼", f"{risk.lawsuit_count} 起"),
        ("司法", "被执行", f"{indicators['executed_count']} 条"),
        ("司法", "失信", f"{indicators['dishonesty_count']} 条"),
        ("司法", "重大诉讼", "是" if indicators["major_lawsuit"] else "否"),
        ("经营", "经营异常", f"{risk.abnormal_operation_count} 次"),
        ("经营", "行政处罚", f"{risk.administrative_penalty_count} 条"),
        ("经营", "法人变更", "频繁" if indicators["legal_person_change_frequent"] else "正常"),
        ("经营", "环保处罚", f"{indicators['env_penalty_count']} 条"),
        ("资金链", "对外担保", f"{indicators['guarantee_count']} 次"),
        ("资金链", "股权质押", f"{indicators['pledge_count']} 次"),
        ("资金链", "破产/清算", f"{indicators['bankruptcy_count']} 次"),
    ]:
        ws3.cell(row, 1, cat).border = thin_border
        ws3.cell(row, 2, label).border = thin_border
        ws3.cell(row, 3, val).border = thin_border
        row += 1

    ws3.column_dimensions["A"].width = 10
    ws3.column_dimensions["B"].width = 14
    ws3.column_dimensions["C"].width = 14

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_html_report(company_name: str) -> str:
    profile = get_company_profile(company_name)
    risk = get_risk_info(company_name)
    indicators = get_risk_indicators(company_name)
    fin = get_financial_metrics(company_name)

    color = "#16a34a"
    if indicators["dishonesty_count"] > 0 or indicators["major_lawsuit"]:
        color = "#dc2626"
    elif risk.lawsuit_count >= 3 or indicators["executed_count"] > 0:
        color = "#d97706"

    def _row(k, v, w=False):
        style = "color:#dc2626;font-weight:600;" if w else ""
        return f"<tr><td style='padding:4px 12px;'>{k}</td><td style='padding:4px 12px;{style}'>{v}</td></tr>"

    fin_rows = ""
    if fin:
        items = [
            ("营收增长率", f"{fin.revenue_growth*100:.1f}%"),
            ("净利增长率", f"{fin.net_profit_growth*100:.1f}%"),
            ("资产负债率", f"{fin.debt_ratio*100:.1f}%"),
            ("每股现金流", f"¥{fin.cash_flow:.2f}"),
            ("ROE", f"{fin.roe*100:.1f}%" if fin.roe else "—"),
            ("流动比率", f"{fin.current_ratio:.2f}" if fin.current_ratio else "—"),
        ]
        fin_rows = "".join(_row(k, v) for k, v in items)
    else:
        fin_rows = "<tr><td colspan='2' style='color:#999;padding:8px 12px;'>无财报数据（非上市公司）</td></tr>"

    risk_rows = "".join([
        _row("诉讼", f"{risk.lawsuit_count} 起"),
        _row("被执行", f"{indicators['executed_count']} 条", indicators["executed_count"] > 0),
        _row("失信", f"{indicators['dishonesty_count']} 条", indicators["dishonesty_count"] > 0),
        _row("重大诉讼", "⚠️ 是" if indicators["major_lawsuit"] else "否", indicators["major_lawsuit"]),
        _row("经营异常", f"{risk.abnormal_operation_count} 次"),
        _row("行政处罚", f"{risk.administrative_penalty_count} 条"),
        _row("法人频繁变更", "⚠️ 是" if indicators["legal_person_change_frequent"] else "否", indicators["legal_person_change_frequent"]),
        _row("对外担保", f"{indicators['guarantee_count']} 次"),
        _row("股权质押", f"{indicators['pledge_count']} 次"),
        _row("破产/清算", f"{indicators['bankruptcy_count']} 次", indicators["bankruptcy_count"] > 0),
        _row("环保处罚", f"{indicators['env_penalty_count']} 条"),
    ])

    return f"""<!DOCTYPE html>
<html lang="zh">
<head>
<meta charset="UTF-8">
<title>风险评估报告 — {company_name}</title>
<style>
  body {{ font-family: -apple-system, "Noto Sans SC", sans-serif; max-width: 800px; margin: 0 auto; padding: 40px 20px; color: #333; }}
  h1 {{ font-size: 24px; border-bottom: 3px solid {color}; padding-bottom: 12px; }}
  .meta {{ color: #888; font-size: 14px; margin-bottom: 24px; }}
  h2 {{ font-size: 16px; color: {color}; margin-top: 28px; }}
  table {{ width: 100%; border-collapse: collapse; }}
  td {{ border-bottom: 1px solid #e8e8e3; font-size: 14px; }}
  .footer {{ margin-top: 40px; color: #bbb; font-size: 12px; text-align: center; }}
  @media print {{ body {{ padding: 0; }} }}
</style>
</head>
<body>
<h1>🔍 企业风险评估报告</h1>
<p class="meta">
  <strong>{company_name}</strong>  |  法人代表: {profile.legal_person}  |  注册资本: {profile.registered_capital}  |  成立: {profile.establish_time}<br>
  报告生成时间: 自动生成  |  本报告仅供参考
</p>

<h2>📊 财务指标</h2>
<table>{fin_rows}</table>

<h2>⚖️ 风险明细</h2>
<table>{risk_rows}</table>

<div class="footer">供应商风险智能分析系统 · 自动生成</div>
</body>
</html>"""
