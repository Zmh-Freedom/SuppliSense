#!/usr/bin/env python3
"""Convert the legacy supplier registration workbook into Feishu Bitable CSVs.

The converter only reads the source workbook. It writes three import-ready CSVs
and a JSON report describing defaults, skipped rows, and source data that has no
direct field in the three-table contract.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MASTER_HEADERS = [
    "供应商代码", "供应商名称", "供应商简称", "统一社会信用代码", "供应商状态",
    "所属行业", "主要产品", "经营范围", "公司网址", "注册地址省份",
    "注册地址城市", "注册地址", "成立日期", "法定代表人", "是否上市",
    "母公司名称", "母公司统一社会信用代码", "数据更新时间", "备注",
]
CAPABILITY_HEADERS = [
    "供应商代码", "品类", "产品名称", "产品关键词", "工艺能力",
    "是否具备设计开发能力", "供货区域", "生产地", "产能说明", "相关资质概况",
    "能力状态", "更新时间",
]
CONTACT_HEADERS = [
    "供应商代码", "联系人姓名", "职务", "联系人类型", "电话", "邮箱",
    "是否主要联系人", "是否已验证", "验证时间", "备注",
]

EMPTY_MARKERS = {"无", "暂无", "无此项", "-", "—", "/", "n/a", "na"}


def _normalise_header(value: Any) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"[\s*()（）\-_：:]", "", text).casefold()


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).replace("\u00a0", " ").strip()
    return "" if text.casefold() in EMPTY_MARKERS else text


def _value(row: dict[str, Any], *names: str) -> str:
    for name in names:
        value = _clean(row.get(_normalise_header(name)))
        if value:
            return value
    return ""


def _split_values(value: str) -> list[str]:
    return [part.strip() for part in re.split(r"[,，、;/；\n]+", value) if part.strip()]


def _normalise_yes_no(value: str, *, unknown: str = "未知") -> str:
    lowered = value.casefold()
    if lowered in {"是", "yes", "true", "1", "y"}:
        return "是"
    if lowered in {"否", "no", "false", "0", "n"}:
        return "否"
    return unknown


def _parse_date(value: str, *, with_time: bool = False) -> str:
    if not value:
        return ""
    candidates = (
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d", "%Y年%m月%d日",
    )
    for pattern in candidates:
        try:
            parsed = datetime.strptime(value, pattern)
            return parsed.strftime("%Y-%m-%d %H:%M:%S" if with_time else "%Y-%m-%d")
        except ValueError:
            continue
    return value


def _read_rows(workbook: Any, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        return []
    sheet = workbook[sheet_name]
    headers = [_normalise_header(cell.value) for cell in sheet[1]]
    rows: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {header: value for header, value in zip(headers, values) if header}
        if any(_clean(value) for value in row.values()):
            rows.append(row)
    return rows


def _merge_non_empty(old: dict[str, str], new: dict[str, str]) -> dict[str, str]:
    return {key: old.get(key) or value for key, value in new.items()}


def _infer_contact_type(title: str) -> str:
    if any(word in title for word in ("技术", "研发", "工程")):
        return "技术"
    if "质量" in title or "品质" in title:
        return "质量"
    if any(word in title for word in ("财务", "会计")):
        return "财务"
    if any(word in title for word in ("商务", "销售", "采购", "业务")):
        return "商务"
    return "其他"


def _join_location(province: str, city: str) -> str:
    return "/".join(part for part in (province, city) if part)


def _write_csv(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def convert_workbook(
    input_path: Path,
    output_dir: Path,
    *,
    default_status: str = "正常",
    default_category: str = "待分类",
    capability_status: str = "待验证",
    updated_at: str | None = None,
) -> dict[str, Any]:
    """Convert a legacy workbook and return the generated conversion report."""
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    report: dict[str, Any] = {
        "input_file": str(input_path),
        "output_dir": str(output_dir),
        "defaults": {
            "supplier_status": default_status,
            "capability_status": capability_status,
            "category": default_category,
            "updated_at": updated_at or "input_file_mtime",
        },
        "counts": {},
        "warnings": [],
        "errors": [],
        "unmapped_source_sheets": ["5-业务能力", "6-主要客户", "7-关联企业"],
    }

    if not updated_at:
        updated_at = datetime.fromtimestamp(input_path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")

    master_rows: dict[str, dict[str, str]] = {}
    for source in _read_rows(workbook, "1-基础信息"):
        code = _value(source, "供货商代码", "供应商代码")
        name = _value(source, "供应商全称", "供应商名称")
        if not code or not name:
            report["errors"].append({"sheet": "1-基础信息", "reason": "缺少供应商代码或供应商全称"})
            continue
        row = {
            "供应商代码": code,
            "供应商名称": name,
            "供应商简称": _value(source, "供应商简称"),
            "统一社会信用代码": _value(source, "社会统一信用代码", "统一社会信用代码"),
            "供应商状态": default_status,
            "所属行业": "",
            "主要产品": _value(source, "主要产品"),
            "经营范围": _value(source, "经营范围"),
            "公司网址": _value(source, "公司网址"),
            "注册地址省份": _value(source, "注册地址-省份", "注册地址省份"),
            "注册地址城市": _value(source, "注册地址-城市", "注册地址城市"),
            "注册地址": _value(source, "注册地址-详细地址", "注册地址"),
            "成立日期": _parse_date(_value(source, "成立日期")),
            "法定代表人": _value(source, "法定代表人"),
            "是否上市": _normalise_yes_no(_value(source, "是否上市")),
            "母公司名称": _value(source, "母公司全称", "母公司名称"),
            "母公司统一社会信用代码": _value(source, "母公司统一社会信用代码"),
            "数据更新时间": _parse_date(updated_at, with_time=True),
            "备注": "行业字段原表未提供，供应商状态使用脚本默认值",
        }
        if code in master_rows:
            report["warnings"].append({"sheet": "1-基础信息", "supplier_code": code, "reason": "供应商代码重复，已合并非空字段"})
            master_rows[code] = _merge_non_empty(master_rows[code], row)
        else:
            master_rows[code] = row

    masters = list(master_rows.values())
    master_codes = set(master_rows)
    report["counts"]["供应商主数据"] = len(masters)

    qualification_map: dict[str, list[str]] = defaultdict(list)
    for source in _read_rows(workbook, "4-资质证书"):
        code = _value(source, "供应商代码")
        if not code:
            continue
        if code not in master_codes:
            report["warnings"].append({"sheet": "4-资质证书", "supplier_code": code, "reason": "找不到主数据，未关联资质"})
            continue
        qualification = _value(source, "资质名称")
        authority = _value(source, "认证机构")
        number = _value(source, "证书编号")
        start = _parse_date(_value(source, "有效期起"))
        end = _parse_date(_value(source, "有效期止"))
        detail = " / ".join(part for part in (qualification, authority, number) if part)
        if start or end:
            detail += f"（有效期：{start or '未提供'}至{end or '未提供'}）"
        if detail:
            qualification_map[code].append(detail)

    basic_by_code = {row["供应商代码"]: row for row in masters}
    capability_rows: list[dict[str, str]] = []
    production_rows = _read_rows(workbook, "3-生产地信息")
    production_by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for source in production_rows:
        code = _value(source, "供应商代码")
        if code in master_codes:
            production_by_code[code].append(source)
        elif code:
            report["warnings"].append({"sheet": "3-生产地信息", "supplier_code": code, "reason": "找不到主数据，未输出能力记录"})

    for code, master in basic_by_code.items():
        sources = production_by_code.get(code) or [{}]
        for source in sources:
            product = _value(source, "主要产品") or master["主要产品"]
            keywords = ",".join(_split_values(product))
            province = _value(source, "省份") or master["注册地址省份"]
            city = _value(source, "市") or master["注册地址城市"]
            capacity_parts = []
            for label, field in (
                ("厂房面积", "厂房面积（平方米）"),
                ("管理技术类人数", "管理技术类人数"),
                ("生产技能类管理人数", "生产技能类管理人数"),
                ("是否配套SGMW供货", "是否配套SGMW供货"),
            ):
                value = _value(source, field)
                if value:
                    capacity_parts.append(f"{label}：{value}")
            capability_rows.append({
                "供应商代码": code,
                "品类": default_category,
                "产品名称": product,
                "产品关键词": keywords,
                "工艺能力": "",
                "是否具备设计开发能力": _normalise_yes_no(_value(basic_by_code[code], "是否具备设计和开发能力")),
                "供货区域": _join_location(province, city),
                "生产地": _join_location(province, city),
                "产能说明": "；".join(capacity_parts),
                "相关资质概况": "；".join(qualification_map.get(code, [])),
                "能力状态": capability_status,
                "更新时间": _parse_date(updated_at, with_time=True),
            })
    report["counts"]["供应商供货能力"] = len(capability_rows)

    contacts_by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for source in _read_rows(workbook, "2-联系人信息"):
        code = _value(source, "供应商代码")
        if not code or code not in master_codes:
            if code:
                report["warnings"].append({"sheet": "2-联系人信息", "supplier_code": code, "reason": "找不到主数据，未输出联系人"})
            else:
                report["errors"].append({"sheet": "2-联系人信息", "reason": "缺少供应商代码"})
            continue
        contacts_by_code[code].append(source)

    contact_rows: list[dict[str, str]] = []
    for code, sources in contacts_by_code.items():
        preferred = next(
            (index for index, source in enumerate(sources) if _normalise_yes_no(_value(source, "是否商务联系人"), unknown="否") == "是"),
            0,
        )
        for index, source in enumerate(sources):
            title = _value(source, "职务")
            contact_rows.append({
                "供应商代码": code,
                "联系人姓名": _value(source, "联系人姓名"),
                "职务": title,
                "联系人类型": _infer_contact_type(title),
                "电话": _value(source, "电话"),
                "邮箱": _value(source, "邮箱"),
                "是否主要联系人": "是" if index == preferred else "否",
                "是否已验证": "未知",
                "验证时间": "",
                "备注": "是否商务联系人=" + (_value(source, "是否商务联系人") or "未提供"),
            })
    report["counts"]["供应商联系人"] = len(contact_rows)

    _write_csv(output_dir / "供应商主数据.csv", MASTER_HEADERS, masters)
    _write_csv(output_dir / "供应商供货能力.csv", CAPABILITY_HEADERS, capability_rows)
    _write_csv(output_dir / "供应商联系人.csv", CONTACT_HEADERS, contact_rows)
    report["counts"]["warnings"] = len(report["warnings"])
    report["counts"]["errors"] = len(report["errors"])
    with (output_dir / "conversion_report.json").open("w", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2)
    return report


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="将供应商注册 Excel 转换为飞书三表 CSV")
    parser.add_argument("input", type=Path, help="原始供应商注册 Excel 文件路径")
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/feishu_supplier_import"))
    parser.add_argument("--default-status", choices=("正常", "暂停合作", "禁止合作", "已淘汰", "候选"), default="正常")
    parser.add_argument("--default-category", default="待分类", help="原表没有标准品类时使用的占位品类")
    parser.add_argument("--capability-status", choices=("已验证", "待验证", "已失效"), default="待验证")
    parser.add_argument("--updated-at", help="来源数据更新时间，格式 YYYY-MM-DD HH:MM:SS；不提供时使用源文件修改时间")
    return parser


def main() -> int:
    args = _build_parser().parse_args()
    if not args.input.exists():
        raise SystemExit(f"输入文件不存在: {args.input}")
    report = convert_workbook(
        args.input,
        args.output_dir,
        default_status=args.default_status,
        default_category=args.default_category,
        capability_status=args.capability_status,
        updated_at=args.updated_at,
    )
    print(json.dumps({"output_dir": str(args.output_dir), "counts": report["counts"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
